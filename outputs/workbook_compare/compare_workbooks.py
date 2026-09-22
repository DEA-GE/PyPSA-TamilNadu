from __future__ import annotations

import csv
import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
FILE_A = ROOT / "additional_data" / "Capacity_Metatable_Data_Sept26.xlsx"
FILE_B = ROOT / "additional_data" / "TamilNadu_ICED_all_source_1782910007272.xlsx.xlsx"
OUT_DIR = ROOT / "outputs" / "workbook_compare"


def normalize(value):
    if value is None:
        return ""
    return value


def used_bounds(ws):
    cells = [
        (cell.row, cell.column)
        for row in ws.iter_rows()
        for cell in row
        if normalize(cell.value) != ""
    ]
    if not cells:
        return {"min_row": 0, "max_row": 0, "min_col": 0, "max_col": 0, "non_empty": 0}
    rows = [r for r, _ in cells]
    cols = [c for _, c in cells]
    return {
        "min_row": min(rows),
        "max_row": max(rows),
        "min_col": min(cols),
        "max_col": max(cols),
        "non_empty": len(cells),
    }


def workbook_profile(path):
    wb = load_workbook(path, data_only=False, read_only=True)
    try:
        return {
            "path": str(path),
            "sheets": {
                ws.title: {
                    "max_row": ws.max_row,
                    "max_column": ws.max_column,
                    "used_bounds": used_bounds(ws),
                }
                for ws in wb.worksheets
            },
        }
    finally:
        wb.close()


def compare_sheet(ws_a, ws_b):
    max_row = max(ws_a.max_row, ws_b.max_row)
    max_col = max(ws_a.max_column, ws_b.max_column)
    diffs = []
    compared = changed = only_a = only_b = same = 0
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            cell_a = ws_a.cell(row, col)
            cell_b = ws_b.cell(row, col)
            a = normalize(cell_a.value)
            b = normalize(cell_b.value)
            if a == "" and b == "":
                continue
            compared += 1
            coord = cell_a.coordinate
            if a == b:
                same += 1
                continue
            if a == "":
                only_b += 1
                kind = "only_in_file_b"
            elif b == "":
                only_a += 1
                kind = "only_in_file_a"
            else:
                changed += 1
                kind = "changed"
            diffs.append(
                {
                    "cell": coord,
                    "kind": kind,
                    "file_a_value": a,
                    "file_b_value": b,
                }
            )
    return {
        "stats": {
            "compared_non_empty_positions": compared,
            "same": same,
            "changed": changed,
            "only_in_file_a": only_a,
            "only_in_file_b": only_b,
        },
        "diffs": diffs,
    }


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    profile_a = workbook_profile(FILE_A)
    profile_b = workbook_profile(FILE_B)

    names_a = set(profile_a["sheets"])
    names_b = set(profile_b["sheets"])
    common = sorted(names_a & names_b)

    wb_a = load_workbook(FILE_A, data_only=False, read_only=True)
    wb_b = load_workbook(FILE_B, data_only=False, read_only=True)
    try:
        sheet_results = {
            name: compare_sheet(wb_a[name], wb_b[name])
            for name in common
        }
    finally:
        wb_a.close()
        wb_b.close()

    summary = {
        "file_a": profile_a,
        "file_b": profile_b,
        "sheets_only_in_file_a": sorted(names_a - names_b),
        "sheets_only_in_file_b": sorted(names_b - names_a),
        "common_sheets": common,
        "common_sheet_stats": {
            name: result["stats"] for name, result in sheet_results.items()
        },
    }

    (OUT_DIR / "summary.json").write_text(json.dumps(summary, indent=2, default=str), encoding="utf-8")

    with (OUT_DIR / "cell_differences.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=["sheet", "cell", "kind", "file_a_value", "file_b_value"],
        )
        writer.writeheader()
        for sheet, result in sheet_results.items():
            for diff in result["diffs"]:
                writer.writerow({"sheet": sheet, **diff})

    source_map = {
        "bio power": "bio-power",
        "coal": "coal",
        "hydro": "hydro",
        "nuclear": "nuclear",
        "oil & gas": "oil-gas",
        "small-hydro": "small-hydro",
        "solar": "solar",
        "wind": "wind",
    }
    a_df = pd.read_excel(FILE_A)
    b_df = pd.read_excel(FILE_B)
    a_df = a_df[pd.to_numeric(a_df["Capacity in MW"], errors="coerce").notna()].copy()
    a_2026 = (
        a_df[a_df["Year"].astype(str).eq("2026-27")]
        .assign(source_key=lambda df: df["Source"].astype(str).str.strip().str.lower())
        .groupby("source_key", as_index=False)["Capacity in MW"]
        .sum()
    )
    b_df = b_df[pd.to_numeric(b_df["Capacity (MW)"], errors="coerce").notna()].copy()
    b_df["source_key"] = b_df["Source"].astype(str).str.strip().str.lower().map(source_map)
    b_grouped = (
        b_df.groupby(["source_key", "Commissioning Group"], dropna=False)["Capacity (MW)"]
        .sum()
        .unstack(fill_value=0)
        .reset_index()
    )
    for col in ["operational", "Temporarily Closed", "pipeline", "retired"]:
        if col not in b_grouped:
            b_grouped[col] = 0
    semantic = a_2026.merge(b_grouped, on="source_key", how="outer").fillna(0)
    semantic["b_operational_plus_temp_closed"] = semantic["operational"] + semantic["Temporarily Closed"]
    semantic["diff_a_2026_minus_b_operational"] = semantic["Capacity in MW"] - semantic["operational"]
    semantic["diff_a_2026_minus_b_operational_plus_temp_closed"] = (
        semantic["Capacity in MW"] - semantic["b_operational_plus_temp_closed"]
    )
    semantic = semantic[
        [
            "source_key",
            "Capacity in MW",
            "operational",
            "Temporarily Closed",
            "b_operational_plus_temp_closed",
            "pipeline",
            "retired",
            "diff_a_2026_minus_b_operational",
            "diff_a_2026_minus_b_operational_plus_temp_closed",
        ]
    ].sort_values("source_key")
    semantic.to_csv(OUT_DIR / "source_aggregate_comparison.csv", index=False)

    print(json.dumps(summary, indent=2, default=str))
    print(f"Saved cell-level differences to {OUT_DIR / 'cell_differences.csv'}")
    print(f"Saved source aggregate comparison to {OUT_DIR / 'source_aggregate_comparison.csv'}")
    print(semantic.to_string(index=False))


if __name__ == "__main__":
    main()
