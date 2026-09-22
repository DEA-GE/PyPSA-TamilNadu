from __future__ import annotations

import csv
import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
FILE_A = ROOT / "additional_data" / "TamilNadu_ICED_all_source_1782910007272.xlsx.xlsx"
FILE_B = ROOT / "additional_data" / "TamilNadu_plantCapacity_all_source_Sep26.xlsx.xlsx"
OUT_DIR = ROOT / "outputs" / "workbook_compare_sep26"


def norm(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return value


def used_bounds(ws):
    cells = [
        (cell.row, cell.column)
        for row in ws.iter_rows()
        for cell in row
        if norm(cell.value) != ""
    ]
    if not cells:
        return {"min_row": 0, "max_row": 0, "min_col": 0, "max_col": 0, "non_empty": 0}
    return {
        "min_row": min(r for r, _ in cells),
        "max_row": max(r for r, _ in cells),
        "min_col": min(c for _, c in cells),
        "max_col": max(c for _, c in cells),
        "non_empty": len(cells),
    }


def profile(path):
    wb = load_workbook(path, data_only=False, read_only=True)
    try:
        return {
            "path": str(path),
            "sheets": {
                ws.title: {
                    "max_row": ws.max_row,
                    "max_column": ws.max_column,
                    "used_bounds": used_bounds(ws),
                    "headers": [norm(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))],
                }
                for ws in wb.worksheets
            },
        }
    finally:
        wb.close()


def cell_diffs(sheet_name):
    a_df = pd.read_excel(FILE_A, sheet_name=sheet_name, header=None, dtype=object).fillna("")
    b_df = pd.read_excel(FILE_B, sheet_name=sheet_name, header=None, dtype=object).fillna("")
    max_row = max(len(a_df), len(b_df))
    max_col = max(len(a_df.columns), len(b_df.columns))
    a_df = a_df.reindex(index=range(max_row), columns=range(max_col), fill_value="")
    b_df = b_df.reindex(index=range(max_row), columns=range(max_col), fill_value="")
    diffs = []
    stats = {"same": 0, "changed": 0, "only_in_a": 0, "only_in_b": 0}
    for row in range(max_row):
        for col in range(max_col):
            a = norm(a_df.iat[row, col])
            b = norm(b_df.iat[row, col])
            if a == "" and b == "":
                continue
            if a == b:
                stats["same"] += 1
                continue
            kind = "changed"
            if a == "":
                kind = "only_in_b"
            elif b == "":
                kind = "only_in_a"
            stats[kind] += 1
            coord = f"{get_column_letter(col + 1)}{row + 1}"
            diffs.append(
                {
                    "sheet": sheet_name,
                    "cell": coord,
                    "kind": kind,
                    "file_a_value": a,
                    "file_b_value": b,
                }
            )
    return stats, diffs


def row_keyed_diff():
    a = pd.read_excel(FILE_A, dtype=str).fillna("")
    b = pd.read_excel(FILE_B, dtype=str).fillna("")
    common_cols = [c for c in a.columns if c in set(b.columns)]
    key_cols = [
        "Source",
        "Commissioning Group",
        "Name of Power Plant",
        "State",
        "District",
        "Commissioning Date",
        "Capacity (MW)",
        "Implementing Agency",
    ]
    key_cols = [c for c in key_cols if c in common_cols]

    def clean(df):
        out = df[common_cols].copy()
        for col in common_cols:
            out[col] = out[col].astype(str).str.strip()
        return out

    a_clean = clean(a)
    b_clean = clean(b)
    a_keys = set(map(tuple, a_clean[key_cols].to_numpy()))
    b_keys = set(map(tuple, b_clean[key_cols].to_numpy()))
    only_a = sorted(a_keys - b_keys)
    only_b = sorted(b_keys - a_keys)

    agg_cols = ["Source", "Commissioning Group"]
    a_num = pd.read_excel(FILE_A)
    b_num = pd.read_excel(FILE_B)
    agg = (
        a_num.groupby(agg_cols, dropna=False)["Capacity (MW)"].agg(file_a_count="count", file_a_mw="sum")
        .reset_index()
        .merge(
            b_num.groupby(agg_cols, dropna=False)["Capacity (MW)"].agg(file_b_count="count", file_b_mw="sum").reset_index(),
            on=agg_cols,
            how="outer",
        )
        .fillna(0)
    )
    agg["mw_diff_b_minus_a"] = agg["file_b_mw"] - agg["file_a_mw"]
    agg["count_diff_b_minus_a"] = agg["file_b_count"] - agg["file_a_count"]
    return key_cols, only_a, only_b, agg.sort_values(agg_cols)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prof_a = profile(FILE_A)
    prof_b = profile(FILE_B)
    common = sorted(set(prof_a["sheets"]) & set(prof_b["sheets"]))

    all_diffs = []
    sheet_stats = {}
    for sheet in common:
        stats, diffs = cell_diffs(sheet)
        sheet_stats[sheet] = stats
        all_diffs.extend(diffs)

    with (OUT_DIR / "cell_differences.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["sheet", "cell", "kind", "file_a_value", "file_b_value"])
        writer.writeheader()
        writer.writerows(all_diffs)

    key_cols, only_a, only_b, agg = row_keyed_diff()
    pd.DataFrame(only_a, columns=key_cols).to_csv(OUT_DIR / "rows_only_in_original.csv", index=False)
    pd.DataFrame(only_b, columns=key_cols).to_csv(OUT_DIR / "rows_only_in_sep26.csv", index=False)
    agg.to_csv(OUT_DIR / "capacity_by_source_group.csv", index=False)

    summary = {
        "file_a": prof_a,
        "file_b": prof_b,
        "sheets_only_in_file_a": sorted(set(prof_a["sheets"]) - set(prof_b["sheets"])),
        "sheets_only_in_file_b": sorted(set(prof_b["sheets"]) - set(prof_a["sheets"])),
        "common_sheets": common,
        "cell_diff_stats": sheet_stats,
        "row_key_columns": key_cols,
        "rows_only_in_file_a": len(only_a),
        "rows_only_in_file_b": len(only_b),
        "total_cell_differences": len(all_diffs),
    }
    (OUT_DIR / "summary.json").write_text(json.dumps(summary, indent=2, default=str), encoding="utf-8")
    print(json.dumps(summary, indent=2, default=str))
    print("\nCapacity by Source / Commissioning Group:")
    print(agg.to_string(index=False))
    print(f"\nSaved outputs in {OUT_DIR}")


if __name__ == "__main__":
    main()
